"""Perfilado local de CSV/XLSX; no utiliza ningún proveedor de IA.

Etapa 1: exploración y diagnóstico, sin modificar ni persistir el archivo.
"""

from __future__ import annotations

import csv
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile, is_zipfile

import pandas as pd
from openpyxl import load_workbook

from app.services.column_semantics import classify_column, suspected_missing_mask
from app.services.measure_hints import measure_hints

MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES = 60 * 1024 * 1024
MAX_XLSX_ENTRIES = 1000
MAX_COLUMNS = 250
MAX_ROWS = 250_000
# Límite de celdas (filas × columnas): acota la memoria aunque el archivo pese
# menos de 25 MB (por ejemplo, millones de celdas de un carácter). SIACAM tiene
# ≈ 0,8 millones de celdas.
MAX_CELLS = 5_000_000


def _shape_error(max_cells: int) -> "DatasetProfileError":
    cells = f"{max_cells:,}".replace(",", ".")
    rows = f"{MAX_ROWS:,}".replace(",", ".")
    return DatasetProfileError(
        f"El dataset supera los límites del MVP: {rows} filas, {MAX_COLUMNS} columnas o {cells} celdas."
    )


class DatasetProfileError(ValueError):
    """Error de validación que puede comunicarse al usuario."""


def _read_csv(content: bytes, max_cells: int) -> tuple[pd.DataFrame, str, str | None]:
    encoding = None
    for candidate in ("utf-8-sig", "cp1252"):
        try:
            decoded = content.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue

    if encoding is None:
        raise DatasetProfileError("No se pudo decodificar el CSV.")

    if "\x00" in decoded:
        raise DatasetProfileError("El CSV contiene caracteres nulos no admitidos.")

    try:
        dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","

    _check_csv_shape(decoded, delimiter, max_cells)

    try:
        df = pd.read_csv(
            StringIO(decoded),
            sep=delimiter,
            dtype=str,
            keep_default_na=False,
            on_bad_lines="error",
            low_memory=False,
        )
    except (pd.errors.ParserError, pd.errors.EmptyDataError, ValueError) as exc:
        raise DatasetProfileError("El archivo CSV no tiene una estructura tabular válida.") from exc

    return df, encoding, None


def _check_csv_shape(decoded: str, delimiter: str, max_cells: int) -> None:
    """Valida filas, columnas y celdas ANTES de crear el DataFrame.

    Recorre el texto con `csv.reader` (incremental, sin guardar filas) y corta
    apenas se supera un límite. Si la cantidad de saltos de línea ya garantiza
    que se cumplen los límites, no hace falta recorrerlo.
    """
    reader = csv.reader(StringIO(decoded), delimiter=delimiter)
    try:
        header = next(reader, None)
        if header is None:
            return  # vacío: lo informa la lectura con Pandas
        columns = max(len(header), 1)
        if columns > MAX_COLUMNS:
            raise _shape_error(max_cells)
        line_bound = decoded.count("\n")
        if line_bound <= MAX_ROWS and line_bound * columns <= max_cells:
            return
        rows = 0
        for row in reader:
            rows += 1
            if rows > MAX_ROWS or rows * max(columns, len(row)) > max_cells:
                raise _shape_error(max_cells)
    except csv.Error as exc:
        raise DatasetProfileError("El archivo CSV no tiene una estructura tabular válida.") from exc


def _check_xlsx_archive(content: bytes) -> None:
    """Controla el ZIP del XLSX descomprimiendo de verdad, con tope.

    Los tamaños declarados en el ZIP pueden falsearse; por eso se cuentan los
    bytes reales al descomprimir y se corta al superar el límite.
    """
    if not is_zipfile(BytesIO(content)):
        raise DatasetProfileError("El archivo no es un XLSX válido.")
    try:
        with ZipFile(BytesIO(content)) as zf:
            entries = zf.infolist()
            if len(entries) > MAX_XLSX_ENTRIES:
                raise DatasetProfileError("El XLSX contiene demasiados componentes internos.")
            total = 0
            for info in entries:
                with zf.open(info) as member:
                    while chunk := member.read(1024 * 1024):
                        total += len(chunk)
                        if total > MAX_XLSX_UNCOMPRESSED_BYTES:
                            raise DatasetProfileError("El XLSX descomprimido supera el límite de seguridad.")
    except DatasetProfileError:
        raise
    except Exception as exc:  # ZIP dañado o con métodos no soportados
        raise DatasetProfileError("El archivo XLSX está dañado.") from exc


def _check_xlsx_shape(content: bytes, max_cells: int) -> None:
    """Cuenta filas y columnas de la primera hoja en modo de solo lectura (por flujo)."""
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        raise DatasetProfileError("No se pudo leer el contenido del XLSX.") from exc
    try:
        if not workbook.worksheets:
            raise DatasetProfileError("El XLSX no contiene hojas.")
        rows, width = 0, 0
        for row in workbook.worksheets[0].iter_rows(values_only=True):
            rows += 1
            width = max(width, len(row))
            # La primera fila es el encabezado.
            if width > MAX_COLUMNS or rows - 1 > MAX_ROWS or rows * width > max_cells:
                raise _shape_error(max_cells)
    except DatasetProfileError:
        raise
    except Exception as exc:
        raise DatasetProfileError("No se pudo leer el contenido del XLSX.") from exc
    finally:
        workbook.close()


def _read_excel(content: bytes, max_cells: int) -> tuple[pd.DataFrame, str, str | None]:
    _check_xlsx_archive(content)
    _check_xlsx_shape(content, max_cells)

    try:
        with pd.ExcelFile(BytesIO(content), engine="openpyxl") as workbook:
            if not workbook.sheet_names:
                raise DatasetProfileError("El XLSX no contiene hojas.")
            sheet_name = workbook.sheet_names[0]
            df = workbook.parse(sheet_name=sheet_name, dtype=str, keep_default_na=False)
    except DatasetProfileError:
        raise
    except (ValueError, OSError, BadZipFile, KeyError, TypeError) as exc:
        raise DatasetProfileError("No se pudo leer el contenido del XLSX.") from exc

    return df, "xlsx", sheet_name


def read_dataset(
    content: bytes, filename: str, *, max_bytes: int | None = None, max_cells: int | None = None
) -> tuple[pd.DataFrame, str, str, str | None]:
    """Lee un CSV/XLSX como texto, con todos los controles de tamaño y forma.

    Devuelve (DataFrame, extensión, codificación, hoja). Todas las celdas son
    texto ('' = vacía): así no se convierten identificadores como '00123' ni se
    interpretan marcadores como 'NA'. La usan el perfilador y el motor de
    análisis (`analysis_engine.py`), para que ambos lean exactamente igual.

    `max_bytes` y `max_cells` permiten límites MÁS BAJOS (por ejemplo, en la beta
    pública); nunca superan los del MVP.
    """
    byte_limit = MAX_FILE_BYTES if max_bytes is None else min(max_bytes, MAX_FILE_BYTES)
    cell_limit = MAX_CELLS if max_cells is None else min(max_cells, MAX_CELLS)
    name = Path(filename or "").name
    extension = Path(name).suffix.lower()
    if extension not in (".csv", ".xlsx"):
        raise DatasetProfileError("Solo se permiten archivos CSV y XLSX.")
    if not content:
        raise DatasetProfileError("El archivo está vacío.")
    if len(content) > byte_limit:
        raise DatasetProfileError(f"El archivo supera el límite de {byte_limit // (1024 * 1024)} MB.")

    if extension == ".csv":
        df, encoding, sheet_name = _read_csv(content, cell_limit)
    else:
        df, encoding, sheet_name = _read_excel(content, cell_limit)

    if df.empty or not len(df.columns):
        raise DatasetProfileError("No se encontraron filas y columnas de datos.")
    # Segunda verificación sobre el DataFrame ya leído (defensa en profundidad).
    if df.shape[1] > MAX_COLUMNS or df.shape[0] > MAX_ROWS or df.shape[0] * df.shape[1] > cell_limit:
        raise _shape_error(cell_limit)

    df.columns = [str(column) for column in df.columns]
    return df.fillna(""), extension, encoding, sheet_name


def profile_dataset(content: bytes, filename: str, *, max_bytes: int | None = None, max_cells: int | None = None) -> dict:
    """Lee un dataset tabular y devuelve un diagnóstico JSON serializable.

    No elimina duplicados, rellena nulos ni altera el archivo original.
    La lectura como texto evita convertir automáticamente identificadores '00123'.

    Política de faltantes: solo la celda vacía (o con espacios) es faltante.
    Marcadores como 'NA', 'null' o '-' se informan como *sospechosos* y no se
    descuentan, porque pueden ser valores válidos.
    """
    name = Path(filename or "").name
    df, extension, encoding, sheet_name = read_dataset(content, filename, max_bytes=max_bytes, max_cells=max_cells)
    column_profiles = []
    warnings = []
    for column in df.columns:
        series = df[column].astype(str)
        missing_mask = series.str.strip().eq("")
        missing = int(missing_mask.sum())
        values = series.loc[~missing_mask]
        suspected_mask = suspected_missing_mask(values)
        suspected = int(suspected_mask.sum())
        suspected_values = values.loc[suspected_mask].str.strip().value_counts().head(5)
        semantics = classify_column(column, values.loc[~suspected_mask])
        column_profiles.append({
            "name": column,
            "missing": missing,
            "non_empty": int(len(df) - missing),
            "distinct": int(values.nunique()),
            "examples": values.drop_duplicates().head(3).tolist(),
            "suspected_missing": suspected,
            "suspected_missing_values": {str(k): int(v) for k, v in suspected_values.items()},
            "value_kind": semantics.value_kind,
            "semantic_type": semantics.semantic_type,
            "semantic_confidence": semantics.confidence,
            "semantic_reason": semantics.reason,
        })
        if missing:
            warnings.append(f"La columna '{column}' tiene {missing} valores vacíos.")
        if suspected:
            markers = ", ".join(f"'{k}'" for k in suspected_values.index)
            warnings.append(
                f"La columna '{column}' tiene {suspected} valores como {markers} que podrían indicar "
                "ausencia; no se cuentan como faltantes porque pueden ser valores válidos."
            )
        if semantics.ambiguity:
            warnings.append(f"La columna '{column}' {semantics.ambiguity}")
        if column.startswith("Unnamed:"):
            warnings.append(f"La columna '{column}' no tiene un encabezado claro.")

    # Sugerencias sobre las medidas (unidad, aditividad, dependencia, marcadores numéricos).
    hints = measure_hints(df, {c["name"]: c["semantic_type"] for c in column_profiles})
    for column in column_profiles:
        h = hints.get(column["name"])
        column.update({
            "unit_suggestion": h.unit if h else None,
            "unit_reason": h.unit_reason if h else None,
            "additivity_suggestion": h.additivity if h else None,
            "additivity_reason": h.additivity_reason if h else None,
            "split_by_suggestion": h.split_by if h else [],
            "split_reason": h.split_reason if h else None,
            "missing_value_suggestions": h.missing_values if h else [],
        })
        if h and h.split_by:
            warnings.append(f"La columna '{column['name']}': {h.split_reason}")
        if h and h.missing_values:
            values = ", ".join(f"'{v}'" for v in h.missing_values)
            warnings.append(
                f"La columna '{column['name']}' tiene valores como {values}, fuera del rango del resto: podrían indicar ausencia."
            )

    duplicate_rows = int(df.duplicated().sum())
    if duplicate_rows:
        warnings.append(
            f"Se encontraron {duplicate_rows} filas idénticas. No se eliminan automáticamente."
        )
    if encoding == "cp1252":
        warnings.append("CSV leído como Windows-1252; revisar la codificación de origen.")

    total_cells = int(df.shape[0] * df.shape[1])
    empty_cells = sum(column["missing"] for column in column_profiles)
    suspected_cells = sum(column["suspected_missing"] for column in column_profiles)

    return {
        "filename": name,
        "format": extension.lstrip("."),
        "encoding": encoding if extension == ".csv" else None,
        "sheet_name": sheet_name,
        "row_count": int(df.shape[0]),
        "column_count": int(df.shape[1]),
        "duplicate_rows": duplicate_rows,
        # Completitud sobre el archivo completo: celdas no vacías / total de celdas.
        "total_cells": total_cells,
        "empty_cells": empty_cells,
        "suspected_missing_cells": suspected_cells,
        "completeness": round(1 - empty_cells / total_cells, 6) if total_cells else None,
        "columns": column_profiles,
        "preview": [
            {column: str(value) for column, value in row.items()}
            for row in df.head(5).to_dict(orient="records")
        ],
        "warnings": warnings,
    }
